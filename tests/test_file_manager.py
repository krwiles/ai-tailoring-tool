from datetime import date
from pathlib import Path

import pytest
from docx import Document

from src.model import JobData


# ---- sanitize_dirname ----

@pytest.mark.parametrize(
    "name, expected",
    [
        ("Acme Corp", "Acme Corp"),
        ("Acme/Corp:Inc*?", "Acme_Corp_Inc__"),
        ("  spaced   out  ", "spaced out"),
        ("trailing dots..  ", "trailing dots"),
        ("", "untitled"),
    ],
)
def test_sanitize_dirname(file_manager, name, expected):
    assert file_manager.sanitize_dirname(name) == expected


# ---- load_text ----

def test_load_text_reads_file(file_manager, tmp_path):
    (tmp_path / "prompts").mkdir()
    (tmp_path / "prompts" / "note.txt").write_text("hello world", encoding="utf-8")

    assert file_manager.load_text(Path("prompts", "note.txt")) == "hello world"


def test_load_text_missing_file_raises(file_manager):
    with pytest.raises(FileNotFoundError):
        file_manager.load_text(Path("nope.txt"))


# ---- read_docx / write_docx ----

def test_read_docx_joins_paragraphs(file_manager, make_docx):
    path = make_docx("data/note.docx", ["Line one", "Line two"])
    assert file_manager.read_docx(path) == "Line one\nLine two"


def test_write_docx_round_trips(file_manager, fake_settings):
    file_manager.write_docx("Line A\nLine B", Path("out.docx"))

    doc = Document(fake_settings.output_dir / "out.docx")
    assert [p.text for p in doc.paragraphs] == ["Line A", "Line B"]


# ---- list_files ----

def test_list_files_filters_and_sorts(file_manager, tmp_path):
    folder = tmp_path / "data" / "resumes"
    folder.mkdir(parents=True)
    for name in ["b.docx", "A.docx", ".hidden.docx", "~$lock.docx", "notes.txt"]:
        (folder / name).write_text("x")

    result = file_manager.list_files(Path("data", "resumes"), "*.docx")

    assert result == [
        Path("data", "resumes", "A.docx"),
        Path("data", "resumes", "b.docx"),
    ]


def test_list_files_missing_directory_returns_empty(file_manager):
    assert file_manager.list_files(Path("does", "not", "exist"), "*.docx") == []


# ---- copy_resume ----

def test_copy_resume_copies_to_destination(file_manager, fake_settings, make_docx):
    template = make_docx("data/resumes/template.docx", ["Resume content"])

    file_manager.copy_resume(Path("CompanyX", "RoleY"), template)

    dest = fake_settings.output_dir / "CompanyX" / "RoleY" / fake_settings.resume_name
    assert dest.exists()
    assert [p.text for p in Document(dest).paragraphs] == ["Resume content"]


def test_copy_resume_missing_template_raises(file_manager):
    with pytest.raises(FileNotFoundError):
        file_manager.copy_resume(Path("CompanyX"), Path("data", "resumes", "missing.docx"))


# ---- copy_cover_letter ----

def test_copy_cover_letter_substitutes_placeholders(file_manager, fake_settings, make_docx):
    template = make_docx(
        "data/cover_letters/template.docx",
        ["Dear {company}, applying for {position} in {location} on {date}."],
    )
    job = JobData(
        job_title="Engineer",
        company="Acme",
        location="Remote",
        job_description="n/a",
    )

    file_manager.copy_cover_letter(Path("Acme", "Engineer"), job, template)

    dest = fake_settings.output_dir / "Acme" / "Engineer" / fake_settings.cover_letter_name
    today = date.today().strftime("%B %d, %Y")
    expected = f"Dear Acme, applying for Engineer in Remote on {today}."
    assert [p.text for p in Document(dest).paragraphs] == [expected]
